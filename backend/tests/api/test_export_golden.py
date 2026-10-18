"""Golden snapshots of every table export.

One deterministic project holds every case the export rules distinguish:
images and videos, boxes on and off a video's saved frame, verified,
rejected and below-threshold boxes, an excluded species, empty files, a
file without a capture date, events with counts, a deployment without a
site, and stored EXIF. Every table export of it is written the way users
get it and compared, text for text, with the snapshots in
``tests/data/export_golden/``.

The point is to change how the tables are produced without changing
what they say. A refactor that keeps the snapshots identical kept the
contents identical; a change that is meant to move a number shows up as
a diff in the snapshot, reviewed with the code.

Regenerate after an intended change with ``ADDAXAI_UPDATE_GOLDEN=1``,
then read the diff before committing it.
"""

from __future__ import annotations

import io
import json
import os
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.api.crud import export as export_crud
from app.ml.postprocessing_outputs import run_readme
from app.ml.postprocessing_outputs.tables_csv import write_tables_csv
from app.ml.postprocessing_outputs.tables_xlsx import write_tables_xlsx
from app.models import Event, EventObservation, LabelTaxonomy, event_files
from tests.api.test_export import _run_camtrap_dp_export
from tests.conftest import (
    make_deployment,
    make_detection,
    make_file,
    make_project,
    make_site,
)

GOLDEN_DIR = Path(__file__).resolve().parent.parent / "data" / "export_golden"
UPDATE = os.environ.get("ADDAXAI_UPDATE_GOLDEN") == "1"

MODEL = "GOLDEN-MODEL"


def _id(name: str) -> str:
    """Readable, stable ids, so a snapshot diff names what moved."""
    return f"{name}-0000"


def _taxonomy(db, name: str, level: str = "species", **ranks: str) -> LabelTaxonomy:
    row = LabelTaxonomy(
        id=_id(f"tax-{name.replace(' ', '-')}"),
        classification_model_id=MODEL,
        name=name,
        level=level,
        **ranks,
    )
    db.add(row)
    db.flush()
    return row


def _event(db, name: str, deployment_id: str, start: datetime, files: list) -> Event:
    ev = Event(
        id=_id(f"ev-{name}"),
        deployment_id=deployment_id,
        event_start_local=start,
        event_end_local=start,
        file_count=len(files),
    )
    db.add(ev)
    db.flush()
    for seq, f in enumerate(files):
        db.execute(
            event_files.insert().values(event_id=ev.id, file_id=f.id, sequence_number=seq)
        )
    return ev


def _build_golden_project(db):
    project = make_project(
        db,
        id=_id("project"),
        name="Golden project",
        timezone="Europe/Amsterdam",
        counting_threshold=0.5,
        detection_model_id="MD-GOLDEN",
        classification_model_id=MODEL,
        excluded_classes=["rat"],
    )
    north = make_site(
        db, project_id=project.id, id=_id("site-north"), name="north",
        latitude=52.1, longitude=5.1, tags={"habitat": "forest"},
    )
    south = make_site(
        db, project_id=project.id, id=_id("site-south"), name="south",
        latitude=51.9, longitude=4.9,
    )
    dep_a = make_deployment(
        db, site_id=north.id, id=_id("dep-a"), folder_path="/data/cam-a",
        start_date_local=date(2024, 6, 1), end_date_local=date(2024, 6, 10),
        camera_model="Cam-X", camera_serial="SN-1", tags={"bait": "none"},
    )
    dep_b = make_deployment(
        db, site_id=south.id, id=_id("dep-b"), folder_path="/data/cam-b",
        start_date_local=date(2024, 6, 2), end_date_local=date(2024, 6, 12),
        paired_cameras=True, notes="two cameras facing",
    )
    dep_c = make_deployment(
        db, project_id=project.id, id=_id("dep-c-nosite"), folder_path="/data/cam-c",
        start_date_local=date(2024, 6, 3), end_date_local=date(2024, 6, 9),
    )

    fox = _taxonomy(
        db, "red fox", taxon_class="mammalia", taxon_order="carnivora",
        taxon_family="canidae", taxon_genus="vulpes", taxon_species="vulpes",
        common_name="Red fox",
    )
    rat = _taxonomy(
        db, "rat", taxon_class="mammalia", taxon_order="rodentia",
        taxon_family="muridae", taxon_genus="rattus", taxon_species="rattus",
    )
    bird = _taxonomy(db, "bird", level="class", taxon_class="aves")
    fox_juv = _taxonomy(
        db, "red fox juvenile", level="variant", taxon_class="mammalia",
        taxon_order="carnivora", taxon_family="canidae", taxon_genus="vulpes",
        taxon_species="vulpes", taxon_variant="juvenile",
    )

    def image(name, dep, when, **kw):
        return make_file(
            db, id=_id(f"img-{name}"), deployment_id=dep.id,
            file_path=f"{dep.folder_path}/{name}.jpg", captured_at_local=when,
            width_px=1920, height_px=1080, **kw,
        )

    def video(name, dep, when, best, **kw):
        return make_file(
            db, id=_id(f"vid-{name}"), deployment_id=dep.id,
            file_path=f"{dep.folder_path}/{name}.mp4", file_type="video",
            file_format="mp4", captured_at_local=when, best_frame_number=best,
            width_px=1280, height_px=720, **kw,
        )

    def box(name, f, conf, **kw):
        return make_detection(db, id=_id(f"det-{name}"), file_id=f.id, confidence=conf, **kw)

    def fox_box(name, f, conf, **kw):
        return box(
            name, f, conf, label="red fox", label_confidence=0.88,
            original_label="red fox", original_label_confidence=0.88,
            label_taxonomy_id=fox.id, scientific_name="V. vulpes",
            common_name="Red fox", classification_method="machine", **kw,
        )

    t = datetime(2024, 6, 3, 21, 15, 0)

    # Deployment A: images, an event with counts, EXIF, the excluded species.
    fox_img = image(
        "fox", dep_a, t,
        exif_data={"Make": "Reconyx\x00", "Model": "HC500", "AmbientTemperature": 14},
        favorited=True,
    )
    fox_box("fox-strong", fox_img, 0.93, bbox_x=0.11, bbox_y=0.21)
    fox_box("fox-weak", fox_img, 0.31)  # below threshold, unverified
    fox_img2 = image("fox-2", dep_a, t.replace(minute=16))
    fox_box("fox-2", fox_img2, 0.81)
    box(
        "fox-2-person", fox_img2, 0.77, category="person", label="person",
        scientific_name="Person", common_name="Person",
    )
    ev_fox = _event(db, "fox", dep_a.id, t, [fox_img, fox_img2])
    db.add(EventObservation(
        id=_id("obs-fox"), event_id=ev_fox.id, label="red fox",
        label_taxonomy_id=fox.id, category="animal", max_n=1, human_count=2,
        sex="female",
    ))
    db.add(EventObservation(
        id=_id("obs-fox-person"), event_id=ev_fox.id, label=None,
        category="person", max_n=1,
    ))
    ev_fox.notes = "two foxes"
    ev_fox.confirmed = True

    rat_img = image("rat-only", dep_a, t.replace(hour=23))
    box(
        "rat", rat_img, 0.9, label="rat", label_confidence=0.9,
        label_taxonomy_id=rat.id, scientific_name="R. rattus",
        common_name="Black rat", classification_method="machine",
    )
    mixed_img = image("rat-and-bird", dep_a, t.replace(hour=23, minute=30))
    box(
        "rat-2", mixed_img, 0.95, label="rat", label_confidence=0.7,
        label_taxonomy_id=rat.id, classification_method="machine",
    )
    box(
        "bird", mixed_img, 0.66, label="bird", label_confidence=0.6,
        label_taxonomy_id=bird.id, scientific_name="Aves", common_name="Bird",
        classification_method="machine",
    )
    verified_rat = image("rat-verified", dep_a, t.replace(day=4))
    box(
        "rat-verified", verified_rat, 0.6, label="rat", label_confidence=1.0,
        label_taxonomy_id=rat.id, classification_method="human", verified=True,
    )

    image("empty", dep_a, t.replace(day=5), verified=True)
    image("no-date", dep_a, None).captured_at_local = None
    rejected = image("rejected", dep_a, t.replace(day=6))
    box(
        "rejected", rejected, 0.7, label="false detection", verified=True,
        classification_method="human",
    )
    low_verified = image("low-verified", dep_a, t.replace(day=6, hour=22))
    fox_box("low-verified", low_verified, 0.12, verified=True)

    # Deployment B: videos with boxes on and off the saved frame.
    vid = video("walk", dep_b, t.replace(day=7), best=3)
    fox_box("vid-f1", vid, 0.7, frame_number=1)
    fox_box("vid-f3", vid, 0.9, frame_number=3)
    fox_box("vid-f3b", vid, 0.6, frame_number=3, bbox_x=0.5)
    fox_box("vid-f5", vid, 0.8, frame_number=5)
    box(
        "vid-f6-verified", vid, 0.4, frame_number=6, label="red fox juvenile",
        label_confidence=1.0, label_taxonomy_id=fox_juv.id,
        classification_method="human", verified=True,
    )
    vid_empty_best = video("empty-best", dep_b, t.replace(day=8), best=2)
    fox_box("vid2-f4", vid_empty_best, 0.85, frame_number=4)
    box("vid2-f0-vehicle", vid_empty_best, 0.9, category="vehicle", frame_number=0)
    vid_no_best = video("no-best", dep_b, t.replace(day=9), best=None)
    fox_box("vid3-f2", vid_no_best, 0.9, frame_number=2)
    ev_vid = _event(db, "walk", dep_b.id, t.replace(day=7), [vid])
    db.add(EventObservation(
        id=_id("obs-walk"), event_id=ev_vid.id, label="red fox",
        label_taxonomy_id=fox.id, category="animal", max_n=2,
    ))

    # Deployment C: no site.
    nosite = image("nosite", dep_c, t.replace(day=4))
    fox_box("nosite", nosite, 0.9)

    for f, obs_type in (
        (fox_img, "animal"), (fox_img2, "animal"), (rat_img, "animal"),
        (mixed_img, "animal"), (verified_rat, "animal"), (low_verified, "animal"),
        (vid, "animal"), (vid_empty_best, "blank"), (vid_no_best, "blank"),
        (nosite, "animal"),
    ):
        f.observation_type = obs_type
    db.commit()
    return project


def _csv_text(data: bytes) -> str:
    return data.decode("utf-8")


def _xlsx_text(data: bytes) -> str:
    """Sheet by sheet, cell values tab-separated: the workbook's content
    without its zip timestamps."""
    wb = load_workbook(io.BytesIO(data), read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            out.append("\t".join("" if v is None else repr(v) for v in row))
    return "\n".join(out) + "\n"


def _camtrap_text(data: bytes) -> str:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for name in sorted(zf.namelist()):
            raw = zf.read(name)
            if name == "datapackage.json":
                package = json.loads(raw)
                package["id"] = "<uuid>"
                package["created"] = "<now>"
                raw = json.dumps(package, indent=2, sort_keys=True).encode()
            out.append(f"## {name}\n{raw.decode('utf-8')}")
    return "\n".join(out)


def _collect_outputs(client, db, project) -> dict[str, str]:
    outputs: dict[str, str] = {}

    for fmt in ("csv", "tsv"):
        folder = tempfile.mkdtemp()
        resp = client.post(
            f"/api/projects/{project.id}/export/tables",
            json={"output_dir": folder, "format": fmt},
        )
        assert resp.status_code == 200, resp.text
        for name in resp.json()["files"]:
            kind = name.split("-", 1)[0]
            outputs[f"tables-{kind}.{fmt}"] = _csv_text((Path(folder) / name).read_bytes())

    folder = tempfile.mkdtemp()
    resp = client.post(
        f"/api/projects/{project.id}/export/tables",
        json={"output_dir": folder, "format": "xlsx"},
    )
    assert resp.status_code == 200, resp.text
    outputs["tables.xlsx.txt"] = _xlsx_text((Path(folder) / resp.json()["files"][0]).read_bytes())

    folder = tempfile.mkdtemp()
    resp = client.post(
        f"/api/projects/{project.id}/export/spatial",
        json={"output_dir": folder, "format": "geojson"},
    )
    assert resp.status_code == 200, resp.text
    outputs["spatial.geojson"] = _csv_text((Path(folder) / resp.json()["files"][0]).read_bytes())

    outputs["camtrap-dp.txt"] = _camtrap_text(
        _run_camtrap_dp_export(client, db, project.id, include_thumbnails=True).content
    )

    folder = Path(tempfile.mkdtemp())
    write_tables_csv(db, project.id, folder)
    for path in sorted(folder.iterdir()):
        outputs[f"folder-run-{path.name}"] = _csv_text(path.read_bytes())
    write_tables_xlsx(db, project.id, folder)
    outputs["folder-run-spreadsheet.xlsx.txt"] = _xlsx_text(
        (folder / "addaxai-spreadsheet.xlsx").read_bytes()
    )

    outputs["readme-counts.json"] = json.dumps(
        {
            "by_category": run_readme._detection_counts_by_category(
                db, project.id, project.counting_threshold
            ),
            "top_species": run_readme._top_species(db, project),
        },
        indent=2,
        sort_keys=True,
    ) + "\n"
    return outputs


def test_every_table_export_matches_its_snapshot(client, db):
    project = _build_golden_project(db)
    outputs = _collect_outputs(client, db, project)

    if UPDATE:
        GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
        for old in GOLDEN_DIR.iterdir():
            old.unlink()
        for name, text in outputs.items():
            (GOLDEN_DIR / name).write_text(text, encoding="utf-8", newline="")
        pytest.skip("golden snapshots regenerated")

    expected = {p.name for p in GOLDEN_DIR.iterdir()}
    assert set(outputs) == expected
    for name, text in outputs.items():
        # newline="" keeps the CSVs' \r\n as written.
        with open(GOLDEN_DIR / name, encoding="utf-8", newline="") as f:
            assert text == f.read(), name


def test_the_fixture_is_not_trivially_empty(client, db):
    """A snapshot of empty tables would pass forever. Guard the guard."""
    project = _build_golden_project(db)
    headers, rows = export_crud.build_detection_rows(db, project)
    assert len(list(rows)) >= 10
    assert headers[0] == "detection_id"


@pytest.mark.parametrize("batch", [1, 3])
def test_batch_boundaries_do_not_change_the_output(client, db, monkeypatch, batch):
    """The exports read the database in batches and regroup boxes per file
    as they go. A file whose boxes straddle a batch edge must come out the
    same, so the snapshots must hold at any batch size, also at one row."""
    if UPDATE:
        pytest.skip("snapshots are being regenerated")
    monkeypatch.setattr(export_crud, "STREAM_BATCH", batch)
    project = _build_golden_project(db)
    outputs = _collect_outputs(client, db, project)
    for name, text in outputs.items():
        with open(GOLDEN_DIR / name, encoding="utf-8", newline="") as f:
            assert text == f.read(), name
